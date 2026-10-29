"""
Generación de reportes del período en PDF, Word o Excel.

El reporte incluye: estadística descriptiva de los registros del período y,
si existen, las proyecciones generadas con ESE MISMO rango de fechas (con sus
métricas MAE / RMSE / R²). El archivo se sube a Supabase Storage (bucket
"reportes", debe ser público) y se registra en la tabla `reportes`.
"""

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from uuid import uuid4

import pandas as pd
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, field_validator

from app.auth import get_current_user, supabase_admin, UsuarioActual

router = APIRouter(prefix="/reportes", tags=["reportes"])

BUCKET = "reportes"

# formato -> (content-type, extensión)
FORMATOS = {
    "pdf": ("application/pdf", "pdf"),
    "word": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "excel": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
}

VARIABLES = {
    "temperatura": ("Temperatura", "°C"),
    "humedad": ("Humedad relativa", "%"),
    "velocidad_viento": ("Velocidad del viento", "km/h"),
}
METODOS = {
    "promedio_movil": "Promedio móvil",
    "holt": "Método de Holt",
    "regresion_lineal": "Regresión lineal múltiple",
}

NOTA_METRICAS = (
    "Nota: MAE y RMSE se calculan con los mismos datos usados para ajustar cada método, "
    "por lo que miden el ajuste al histórico y no la capacidad de predecir días futuros. "
    "Las proyecciones son estimaciones estadísticas, no un pronóstico meteorológico."
)


class GenerarReporteRequest(BaseModel):
    periodo_inicio: date
    periodo_fin: date
    formato: str = "pdf"  # "pdf" | "word" | "excel"

    @field_validator("formato")
    @classmethod
    def formato_valido(cls, v):
        if v not in FORMATOS:
            raise ValueError("formato debe ser 'pdf', 'word' o 'excel'")
        return v

    @field_validator("periodo_fin")
    @classmethod
    def fin_no_menor_que_inicio(cls, v, info):
        inicio = info.data.get("periodo_inicio")
        if inicio and v < inicio:
            raise ValueError("periodo_fin no puede ser anterior a periodo_inicio")
        return v


# ----------------------------------------------------------------------------
# Obtención de datos
# ----------------------------------------------------------------------------
def _cargar_registros(inicio: date, fin: date) -> pd.DataFrame:
    """Trae todos los registros del período (Supabase devuelve máx. 1000 por consulta)."""
    filas, desde, paso = [], 0, 1000
    while True:
        r = (
            supabase_admin.table("registros_climaticos")
            .select("fecha,hora,temperatura,humedad,velocidad_viento,fuente")
            .gte("fecha", inicio.isoformat())
            .lte("fecha", fin.isoformat())
            .order("fecha")
            .order("hora")
            .range(desde, desde + paso - 1)
            .execute()
        )
        filas.extend(r.data)
        if len(r.data) < paso:
            break
        desde += paso
    return pd.DataFrame(filas)


def _estadisticas(df: pd.DataFrame) -> list[dict]:
    resultado = []
    for var, (nombre, unidad) in VARIABLES.items():
        serie = pd.to_numeric(df[var], errors="coerce").dropna() if var in df else pd.Series(dtype=float)
        fila = {"variable": f"{nombre} ({unidad})", "n": int(len(serie))}
        if serie.empty:
            fila.update(media=None, minimo=None, maximo=None, desv=None)
        else:
            fila.update(
                media=round(float(serie.mean()), 2),
                minimo=round(float(serie.min()), 2),
                maximo=round(float(serie.max()), 2),
                desv=round(float(serie.std(ddof=1)) if len(serie) > 1 else 0.0, 2),
            )
        resultado.append(fila)
    return resultado


def _cargar_proyecciones(inicio: date, fin: date) -> list[dict]:
    """Proyecciones generadas con exactamente este período (la más reciente por variable y método)."""
    res = (
        supabase_admin.table("proyecciones")
        .select("*")
        .eq("periodo_inicio", inicio.isoformat())
        .eq("periodo_fin", fin.isoformat())
        .execute()
    )
    filas = sorted(res.data, key=lambda r: r.get("created_at") or "", reverse=True)
    vistos, seleccion = set(), []
    for p in filas:
        clave = (p["variable"], p["metodo"])
        if clave in vistos:
            continue
        vistos.add(clave)
        seleccion.append(p)

    orden_var, orden_met = list(VARIABLES), list(METODOS)
    seleccion.sort(
        key=lambda p: (
            orden_var.index(p["variable"]) if p["variable"] in orden_var else 99,
            orden_met.index(p["metodo"]) if p["metodo"] in orden_met else 99,
        )
    )

    for p in seleccion:
        v = (
            supabase_admin.table("proyeccion_valores")
            .select("fecha,valor")
            .eq("proyeccion_id", p["id"])
            .order("fecha")
            .execute()
        )
        p["valores"] = v.data
    return seleccion


# ----------------------------------------------------------------------------
# Utilidades de formato
# ----------------------------------------------------------------------------
def _f(x, dec=2) -> str:
    return "—" if x is None else f"{float(x):.{dec}f}"


def _fecha_es(s: str) -> str:
    s = str(s)[:10]
    return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"


def _nombre_var(clave: str) -> str:
    n, u = VARIABLES.get(clave, (clave, ""))
    return f"{n} ({u})" if u else n


def _unidad(clave: str) -> str:
    return VARIABLES.get(clave, ("", ""))[1]


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------
def _generar_pdf(d: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    marca = colors.HexColor("#0B4F6C")
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("titulo", parent=estilos["Title"], textColor=marca, fontSize=20, alignment=0)
    h2 = ParagraphStyle("h2", parent=estilos["Heading2"], textColor=marca, spaceBefore=14)
    h3 = ParagraphStyle("h3", parent=estilos["Heading4"], textColor=colors.HexColor("#334155"), spaceBefore=8)
    normal = estilos["BodyText"]
    nota = ParagraphStyle("nota", parent=normal, fontSize=8, textColor=colors.HexColor("#475569"))

    def tabla(filas, anchos=None):
        t = Table(filas, colWidths=anchos, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), marca),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return t

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title="Reporte climático — Cochabamba",
    )
    el = [
        Paragraph("Reporte climático — Cochabamba", titulo),
        Paragraph(f"Período: {_fecha_es(d['inicio'])} al {_fecha_es(d['fin'])}", normal),
        Paragraph(f"Generado el {d['generado']} por {d['usuario']}", normal),
        Paragraph(
            f"{d['n_registros']} registros en {d['n_dias']} día(s) con datos. Fuente: Open-Meteo y registros manuales.",
            normal,
        ),
        Paragraph("Estadística descriptiva", h2),
        tabla(
            [["Variable", "Lecturas", "Media", "Mínimo", "Máximo", "Desv. est."]]
            + [[e["variable"], e["n"], _f(e["media"]), _f(e["minimo"]), _f(e["maximo"]), _f(e["desv"])] for e in d["estadisticas"]],
            [5.5 * cm, 2 * cm, 2.2 * cm, 2.2 * cm, 2.2 * cm, 2.4 * cm],
        ),
        Paragraph("Proyecciones del período", h2),
    ]

    if not d["proyecciones"]:
        el.append(
            Paragraph(
                "No hay proyecciones generadas con este mismo rango de fechas. Genera una en el módulo "
                "de Proyecciones usando el mismo período y vuelve a crear el reporte.",
                normal,
            )
        )
    else:
        el.append(Paragraph("Comparación de métodos", h3))
        el.append(
            tabla(
                [["Variable", "Método", "MAE", "RMSE", "R²"]]
                + [
                    [_nombre_var(p["variable"]), METODOS.get(p["metodo"], p["metodo"]), _f(p.get("mae"), 4), _f(p.get("rmse"), 4), _f(p.get("r2"), 4)]
                    for p in d["proyecciones"]
                ],
                [4.8 * cm, 4.8 * cm, 2.4 * cm, 2.4 * cm, 2.4 * cm],
            )
        )
        el.append(Spacer(1, 6))
        el.append(Paragraph(NOTA_METRICAS, nota))
        for p in d["proyecciones"]:
            u = _unidad(p["variable"])
            bloque = [
                Paragraph(f"{_nombre_var(p['variable'])} — {METODOS.get(p['metodo'], p['metodo'])}", h3),
                tabla(
                    [["Fecha", f"Valor proyectado ({u})"]] + [[_fecha_es(v["fecha"]), _f(v["valor"])] for v in p["valores"]],
                    [4 * cm, 5 * cm],
                ),
            ]
            el.append(KeepTogether(bloque) if len(p["valores"]) <= 12 else bloque[0])
            if len(p["valores"]) > 12:
                el.append(bloque[1])

    doc.build(el)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Word
# ----------------------------------------------------------------------------
def _generar_word(d: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Reporte climático — Cochabamba", 0)
    doc.add_paragraph(f"Período: {_fecha_es(d['inicio'])} al {_fecha_es(d['fin'])}")
    doc.add_paragraph(f"Generado el {d['generado']} por {d['usuario']}")
    doc.add_paragraph(
        f"{d['n_registros']} registros en {d['n_dias']} día(s) con datos. Fuente: Open-Meteo y registros manuales."
    )

    def tabla(encabezados, filas):
        t = doc.add_table(rows=1, cols=len(encabezados))
        t.style = "Table Grid"
        for i, h in enumerate(encabezados):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(h)
            run.bold = True
        for fila in filas:
            celdas = t.add_row().cells
            for i, v in enumerate(fila):
                celdas[i].text = str(v)
                if i > 0:
                    celdas[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for fila in t.rows:
            for c in fila.cells:
                for par in c.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9)
        return t

    doc.add_heading("Estadística descriptiva", 1)
    tabla(
        ["Variable", "Lecturas", "Media", "Mínimo", "Máximo", "Desv. est."],
        [[e["variable"], e["n"], _f(e["media"]), _f(e["minimo"]), _f(e["maximo"]), _f(e["desv"])] for e in d["estadisticas"]],
    )

    doc.add_heading("Proyecciones del período", 1)
    if not d["proyecciones"]:
        doc.add_paragraph(
            "No hay proyecciones generadas con este mismo rango de fechas. Genera una en el módulo de "
            "Proyecciones usando el mismo período y vuelve a crear el reporte."
        )
    else:
        doc.add_heading("Comparación de métodos", 2)
        tabla(
            ["Variable", "Método", "MAE", "RMSE", "R²"],
            [
                [_nombre_var(p["variable"]), METODOS.get(p["metodo"], p["metodo"]), _f(p.get("mae"), 4), _f(p.get("rmse"), 4), _f(p.get("r2"), 4)]
                for p in d["proyecciones"]
            ],
        )
        n = doc.add_paragraph(NOTA_METRICAS)
        for run in n.runs:
            run.italic = True
            run.font.size = Pt(8)
        for p in d["proyecciones"]:
            doc.add_heading(f"{_nombre_var(p['variable'])} — {METODOS.get(p['metodo'], p['metodo'])}", 3)
            tabla(
                ["Fecha", f"Valor proyectado ({_unidad(p['variable'])})"],
                [[_fecha_es(v["fecha"]), _f(v["valor"])] for v in p["valores"]],
            )

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
def _generar_excel(d: dict, df: pd.DataFrame) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    relleno = PatternFill("solid", fgColor="0B4F6C")
    blanco = Font(bold=True, color="FFFFFF")

    def encabezado(ws, fila, textos):
        for i, t in enumerate(textos, start=1):
            c = ws.cell(row=fila, column=i, value=t)
            c.fill, c.font = relleno, blanco
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def anchos(ws, valores):
        for i, w in enumerate(valores, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # --- Hoja 1: Resumen ---
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"], ws["A1"].font = "Reporte climático — Cochabamba", Font(bold=True, size=14, color="0B4F6C")
    ws["A2"] = f"Período: {_fecha_es(d['inicio'])} al {_fecha_es(d['fin'])}"
    ws["A3"] = f"Generado el {d['generado']} por {d['usuario']}"
    ws["A4"] = f"{d['n_registros']} registros en {d['n_dias']} día(s) con datos"
    ws["A6"], ws["A6"].font = "Estadística descriptiva", Font(bold=True, size=12)
    encabezado(ws, 7, ["Variable", "Lecturas", "Media", "Mínimo", "Máximo", "Desv. est."])
    for r, e in enumerate(d["estadisticas"], start=8):
        for c, v in enumerate([e["variable"], e["n"], e["media"], e["minimo"], e["maximo"], e["desv"]], start=1):
            ws.cell(row=r, column=c, value=v)
    anchos(ws, [32, 12, 12, 12, 12, 12])

    # --- Hoja 2: Proyecciones ---
    wp = wb.create_sheet("Proyecciones")
    wp["A1"], wp["A1"].font = "Comparación de métodos", Font(bold=True, size=12)
    if not d["proyecciones"]:
        wp["A3"] = "No hay proyecciones generadas con este mismo rango de fechas."
    else:
        encabezado(wp, 3, ["Variable", "Método", "MAE", "RMSE", "R²"])
        fila = 4
        for p in d["proyecciones"]:
            for c, v in enumerate(
                [_nombre_var(p["variable"]), METODOS.get(p["metodo"], p["metodo"]), p.get("mae"), p.get("rmse"), p.get("r2")],
                start=1,
            ):
                wp.cell(row=fila, column=c, value=v)
            fila += 1
        wp.cell(row=fila + 1, column=1, value=NOTA_METRICAS).alignment = Alignment(wrap_text=True, vertical="top")
        wp.merge_cells(start_row=fila + 1, start_column=1, end_row=fila + 1, end_column=5)
        wp.row_dimensions[fila + 1].height = 48
        fila += 3
        wp.cell(row=fila, column=1, value="Valores proyectados").font = Font(bold=True, size=12)
        fila += 1
        encabezado(wp, fila, ["Variable", "Método", "Fecha", "Valor"])
        for p in d["proyecciones"]:
            for v in p["valores"]:
                fila += 1
                wp.cell(row=fila, column=1, value=_nombre_var(p["variable"]))
                wp.cell(row=fila, column=2, value=METODOS.get(p["metodo"], p["metodo"]))
                wp.cell(row=fila, column=3, value=_fecha_es(v["fecha"]))
                wp.cell(row=fila, column=4, value=v["valor"])
    anchos(wp, [32, 28, 14, 14, 12])

    # --- Hoja 3: Registros (datos del período) ---
    wr = wb.create_sheet("Registros")
    columnas = ["fecha", "hora", "temperatura", "humedad", "velocidad_viento", "fuente"]
    encabezado(wr, 1, ["Fecha", "Hora", "Temperatura (°C)", "Humedad (%)", "Viento (km/h)", "Fuente"])
    for fila in df.reindex(columns=columnas).astype(object).where(df.notna(), None).itertuples(index=False):
        wr.append(list(fila))
    wr.freeze_panes = "A2"
    anchos(wr, [12, 10, 18, 14, 14, 14])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Endpoint
# ----------------------------------------------------------------------------
@router.post("/generar")
def generar_reporte(
    datos: GenerarReporteRequest,
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Genera el reporte del período en el formato pedido y devuelve su dirección de descarga."""
    df = _cargar_registros(datos.periodo_inicio, datos.periodo_fin)
    if df.empty:
        raise HTTPException(status_code=404, detail="No hay registros climáticos en ese período.")

    zona = timezone(timedelta(hours=-4))  # Bolivia
    d = {
        "inicio": datos.periodo_inicio.isoformat(),
        "fin": datos.periodo_fin.isoformat(),
        "generado": datetime.now(zona).strftime("%d/%m/%Y %H:%M"),
        "usuario": usuario.email or usuario.id,
        "n_registros": int(len(df)),
        "n_dias": int(df["fecha"].nunique()),
        "estadisticas": _estadisticas(df),
        "proyecciones": _cargar_proyecciones(datos.periodo_inicio, datos.periodo_fin),
    }

    try:
        if datos.formato == "pdf":
            contenido = _generar_pdf(d)
        elif datos.formato == "word":
            contenido = _generar_word(d)
        else:
            contenido = _generar_excel(d, df)
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Falta instalar una librería en el backend ({exc.name}). Ejecuta: pip install -r requirements.txt",
        )

    content_type, ext = FORMATOS[datos.formato]
    ruta = f"{usuario.id}/{datos.periodo_inicio}_{datos.periodo_fin}_{uuid4().hex[:8]}.{ext}"

    try:
        supabase_admin.storage.from_(BUCKET).upload(
            ruta, contenido, {"content-type": content_type, "upsert": "true"}
        )
        archivo_url = str(supabase_admin.storage.from_(BUCKET).get_public_url(ruta)).rstrip("?")
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"No se pudo guardar el archivo en Supabase Storage (bucket '{BUCKET}'): {exc}",
        )

    fila = {
        "tipo": datos.formato,
        "periodo_inicio": datos.periodo_inicio.isoformat(),
        "periodo_fin": datos.periodo_fin.isoformat(),
        "archivo_url": archivo_url,
    }
    try:
        supabase_admin.table("reportes").insert({**fila, "generado_por": usuario.id}).execute()
    except Exception:
        # Por si la tabla no tiene la columna generado_por
        supabase_admin.table("reportes").insert(fila).execute()

    supabase_admin.table("auditoria").insert(
        {
            "usuario_id": usuario.id,
            "accion": "generar_reporte",
            "tabla_afectada": "reportes",
            "detalle": {
                "formato": datos.formato,
                "periodo_inicio": datos.periodo_inicio.isoformat(),
                "periodo_fin": datos.periodo_fin.isoformat(),
            },
        }
    ).execute()

    return {"mensaje": "Reporte generado correctamente.", "archivo_url": archivo_url, "formato": datos.formato}
